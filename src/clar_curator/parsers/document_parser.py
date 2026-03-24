"""
Parse design documents (Word .docx, PDF) into plain text for the AI layer.
"""

from __future__ import annotations

from pathlib import Path


def parse(file_path: str | Path) -> str:
    """
    Extract text from a Word or PDF design document.

    Returns a single string with the document's full text content.
    """
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".docx":
        return _parse_docx(path)
    elif suffix == ".pdf":
        return _parse_pdf(path)
    else:
        raise ValueError(f"Unsupported document format: {suffix}. Expected .docx or .pdf")


def _parse_docx(path: Path) -> str:
    import docx  # python-docx

    doc = docx.Document(path)
    parts: list[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            # Preserve heading style as markdown-like prefix for context
            if para.style.name.startswith("Heading"):
                level = para.style.name.split()[-1]
                parts.append(f"{'#' * int(level)} {para.text}")
            else:
                parts.append(para.text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))

    return "\n".join(parts)


def _parse_pdf(path: Path) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text)

            # Extract tables from PDF pages too
            for table in page.extract_tables():
                for row in table:
                    cells = [str(cell).strip() if cell else "" for cell in row]
                    parts.append(" | ".join(cells))

    return "\n".join(parts)
